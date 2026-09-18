#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《复数 n 次方根可视化工具 实验报告》Word 文档。

用法（需要 python-docx 与 Pillow）：
    python build_report.py            # 生成 实验报告.docx

报告里的代码片段从 complex_roots.py 直接读取，表格里的数值由 complex_roots.py 现算，
因此文档与程序始终一致，不会出现“文档写一套、代码是另一套”的情况。
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

import complex_roots as cr

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(HERE, "report_figures")
OUT_PATH = os.path.join(HERE, "实验报告.docx")
SRC_PATH = os.path.join(HERE, "complex_roots.py")

BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "2F5496"
ALT_FILL = "F2F5FA"
BORDER = "D9D9D9"


# --------------------------------------------------------------------------- #
# 通用排版工具
# --------------------------------------------------------------------------- #
def set_run_font(run, cn="宋体", latin="Times New Roman", size=None, bold=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.32
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (("Heading 1", 15, 16, 8),
                                      ("Heading 2", 12.5, 12, 6),
                                      ("Heading 3", 11, 10, 5)):
        st = doc.styles[name]
        st.font.name = "微软雅黑"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLACK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.2
        ppr = st.element.get_or_add_pPr()
        for bdr in ppr.findall(qn("w:pBdr")):          # 去掉标题自带的下框线
            ppr.remove(bdr)

    title = doc.styles["Title"]
    title.font.name = "微软雅黑"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.paragraph_format.space_after = Pt(4)
    ppr = title.element.get_or_add_pPr()
    for bdr in ppr.findall(qn("w:pBdr")):
        ppr.remove(bdr)

    section = doc.sections[0]
    section.page_width = Cm(21.0)          # A4，中文报告常用纸张
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def para(doc, text, size=10.5, bold=False, cn="宋体", align=None, space_after=6,
         color=None, indent=None, latin="Times New Roman", space_before=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    set_run_font(run, cn=cn, latin=latin, size=size, bold=bold, color=color)
    return p


def h1(doc, text):
    return doc.add_paragraph(text, style="Heading 1")


def h2(doc, text):
    return doc.add_paragraph(text, style="Heading 2")


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def code_block(doc, text, size=8.5):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.left_indent = Cm(0.6)
        run = p.add_run(line if line else " ")
        set_run_font(run, cn="Consolas", latin="Consolas", size=size,
                     color=RGBColor(0x1F, 0x2A, 0x37))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def caption(doc, text):
    return para(doc, text, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                space_after=12, color=GRAY, cn="微软雅黑")


def table_caption(doc, text):
    """表格标题放在表格上方（中文排版习惯）。"""
    return para(doc, text, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                space_after=3, space_before=6, color=GRAY, cn="微软雅黑")


def add_figure(doc, path, caption_text, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(doc, caption_text)


def add_equation(doc, path, scale=0.75):
    """按 300 dpi 折算宽度插入公式图片，避免放大后发虚。"""
    with Image.open(path) as im:
        w_px = im.size[0]
    width_in = min(5.9, w_px / 300.0) * scale
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(path, width=Inches(width_in))


# --------------------------------------------------------------------------- #
# 表格
# --------------------------------------------------------------------------- #
def _set_cell_bg(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        OxmlElement("w:shd"))
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)

    mar = OxmlElement("w:tblCellMar")
    for side, width in (("top", 90), ("left", 120), ("bottom", 90), ("right", 120)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(width))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def add_table(doc, headers, rows, widths, size=9.5, header_size=9.5,
              align_center_cols=None, header_fill=HEADER_FILL):
    total = 15.5 / float(sum(widths))       # 统一缩放到 15.5cm，确保不超出版心
    widths = [w * total for w in widths]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    align_center_cols = align_center_cols or set()

    hdr = table.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))            # 跨页时重复表头
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        _set_cell_bg(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        set_run_font(run, cn="微软雅黑", size=header_size, bold=True,
                     color=RGBColor(0xFF, 0xFF, 0xFF))

    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Cm(widths[i])
            if r % 2 == 1:
                _set_cell_bg(cells[i], ALT_FILL)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            if i in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(text))
            set_run_font(run, cn="宋体", size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number_footer(doc):
    """页脚居中显示“第 X 页”。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld = p.add_run()
    set_run_font(fld, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    fld._r.append(begin)
    fld._r.append(instr)
    fld._r.append(end)
    run2 = p.add_run(" 页")
    set_run_font(run2, size=9, color=GRAY)


# --------------------------------------------------------------------------- #
# 数学公式图片（Times New Roman / Times Italic，带上下标）
# --------------------------------------------------------------------------- #
_FR = {
    "v": (r"C:\Windows\Fonts\timesi.ttf", 1.00, 0.00),     # 斜体变量
    "n": (r"C:\Windows\Fonts\times.ttf", 1.00, 0.00),      # 正体
    "sub": (r"C:\Windows\Fonts\timesi.ttf", 0.66, 0.30),   # 下标（斜体）
    "subn": (r"C:\Windows\Fonts\times.ttf", 0.66, 0.30),
    "sup": (r"C:\Windows\Fonts\times.ttf", 0.66, -0.46),   # 上标（正体）
}


def render_formula(path, runs, px=64, color=(0, 0, 0), bg=(255, 255, 255)):
    """把 [(文本, 类型)] 渲染成一张公式图片，类型见 _FR。"""
    fonts = {k: ImageFont.truetype(v[0], max(10, int(px * v[1]))) for k, v in _FR.items()}
    probe = ImageDraw.Draw(Image.new("RGB", (4, 4)))
    widths = [probe.textlength(t, font=fonts[k]) for t, k in runs]
    pad = int(px * 0.28)
    W = int(sum(widths)) + pad * 2
    H = int(px * 2.0)
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)
    base = int(px * 1.18)
    x = pad
    for (text, kind), w in zip(runs, widths):
        _, _, dy = _FR[kind]
        d.text((x, base + dy * px), text, font=fonts[kind], fill=color, anchor="ls")
        x += w
    # 裁掉多余留白
    bbox = Image.eval(img.convert("L"), lambda v: 255 - v).getbbox()
    if bbox:
        m = int(px * 0.12)
        img = img.crop((max(0, bbox[0] - m), max(0, bbox[1] - m),
                        min(W, bbox[2] + m), min(H, bbox[3] + m)))
    img.save(path)
    return path


# --------------------------------------------------------------------------- #
# 从源码里取代码片段（保证报告与程序一致）
# --------------------------------------------------------------------------- #
def snippet(marker: str, n_lines: int, back: int = 0, strip_blank=True) -> str:
    with open(SRC_PATH, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    try:
        start = next(i for i, l in enumerate(lines) if marker in l)
    except StopIteration:
        return f"（未找到片段：{marker}）"
    start = max(0, start - back)
    out = lines[start:start + n_lines]
    if strip_blank:
        out = [l for l in out if l.strip()]
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# 正文
# --------------------------------------------------------------------------- #
def build():
    os.makedirs(FIG_DIR, exist_ok=True)

    # ---- 公式图片 ----
    f_mul1 = render_formula(os.path.join(FIG_DIR, "eq_mul1.png"), [
        ("z", "v"), ("1", "sub"), ("·z", "v"), ("2", "sub"), (" = ", "n"),
        ("r", "v"), ("1", "sub"), ("r", "v"), ("2", "sub"), ("[cos(", "n"),
        ("θ", "v"), ("1", "sub"), (" + ", "n"), ("θ", "v"), ("2", "sub"), (") + ", "n"),
        ("i", "v"), (" sin(", "n"), ("θ", "v"), ("1", "sub"), (" + ", "n"), ("θ", "v"),
        ("2", "sub"), (")]", "n")])
    f_mul2 = render_formula(os.path.join(FIG_DIR, "eq_mul2.png"), [
        ("z", "v"), ("1", "sub"), ("·z", "v"), ("2", "sub"), (" = ", "n"),
        ("r", "v"), ("1", "sub"), ("r", "v"), ("2", "sub"), (" e", "n"),
        ("i(θ1+θ2)", "sup")])
    f_spiral = render_formula(os.path.join(FIG_DIR, "eq_spiral.png"), [
        ("z", "v"), ("1", "sub"), ("·z", "v"), ("2", "sub"), ("t", "sup"), (" = ", "n"),
        ("r", "v"), ("1", "sub"), ("r", "v"), ("2", "sub"), ("t", "sup"), (" e", "n"),
        ("i(θ1 + tθ2)", "sup")])
    f_demoivre = render_formula(os.path.join(FIG_DIR, "eq_demoivre.png"), [
        ("(cos ", "n"), ("θ", "v"), (" + ", "n"), ("i", "v"), (" sin ", "n"), ("θ", "v"),
        (")", "n"), ("n", "sup"), (" = cos ", "n"), ("nθ", "v"), (" + ", "n"),
        ("i", "v"), (" sin ", "n"), ("nθ", "v")])
    f_root = render_formula(os.path.join(FIG_DIR, "eq_root.png"), [
        ("w", "v"), ("k", "sub"), (" = ", "n"), ("r", "v"), ("1/n", "sup"), (" · e", "n"),
        ("i(θ + 2πk)/n", "sup"), ("      ", "n"), ("k", "v"), (" = 0, 1, ..., ", "n"),
        ("n", "v"), ("-1", "n")])
    f_check = render_formula(os.path.join(FIG_DIR, "eq_check.png"), [
        ("w", "v"), ("k", "sub"), ("n", "sup"), (" = ", "n"), ("z", "v"),
        ("   <=>   |", "n"), ("w", "v"), ("k", "sub"), ("|", "n"), ("n", "sup"),
        (" = |", "n"), ("z", "v"), ("| ,    ", "n"), ("n", "v"), ("·arg ", "n"),
        ("w", "v"), ("k", "sub"), (" = ", "n"), ("θ", "v"), (" + 2π", "n"), ("k", "v")])

    # ---- 报告里用到的数据（全部现算） ----
    z_demo, n_demo = 3 + 4j, 5
    roots, r_demo, th_demo, R_demo = cr.nth_roots_of(z_demo, n_demo)
    w0 = roots[0]
    power_rows = []
    for j in range(1, n_demo + 1):
        v = w0 ** j
        note = "= z" if j == n_demo else ""
        power_rows.append([j, cr.fmt_complex(v, 4), cr.fmt_num(abs(v), 4),
                           f"{cr.math.degrees(cr.cmath.phase(v)):.2f}°", note])

    z1, z2 = 3 + 1j, 2 + 1j
    p_mul = z1 * z2
    mul_rows = [
        ["由 z1 出发", f"模长 × |z2| = ×{cr.fmt_num(abs(z2), 4)}，辐角 + θ2 = +{cr.math.degrees(cr.cmath.phase(z2)):.2f}°",
         cr.fmt_complex(p_mul, 4)],
        ["由 z2 出发", f"模长 × |z1| = ×{cr.fmt_num(abs(z1), 4)}，辐角 + θ1 = +{cr.math.degrees(cr.cmath.phase(z1)):.2f}°",
         cr.fmt_complex(p_mul, 4)],
    ]

    prompt_rows = [
        ["1", "确定技术路线：要动图 + 交互，且本机没装 numpy/matplotlib",
         "“用 Python 做复数 n 次方根的动图，可以输入任意复数和 n。注意：环境里只有标准库和 Pillow，不要用 numpy 和 matplotlib。”",
         "AI 给了 Pillow 逐帧绘制 + tkinter 的方案，符合约束，采纳；同时确定用 GIF 作为动图格式。"],
        ["2", "数学部分必须严格",
         "“核心公式按德莫弗定理：根为 r^(1/n)·e^(i(θ+2πk)/n)，k 从 0 到 n-1；并给出 |w_k^n − z| 的自校验。”",
         "AI 写出 nth_roots_of()，我补上 k 的取值范围和误差校验的输出。"],
        ["3", "复数输入要宽容",
         "“用户可能输入 3+4i、-8、2i、sqrt(2)-sqrt(2)i、2*exp(i*pi/3)，都要能解析；非法输入要给提示不能崩。”",
         "AI 第一版直接用 complex() 转换，只能处理 j 形式；我改为自写词法与语法分析，并补上省略乘号（3i、2(1+i)）。"],
        ["4", "动画要讲清楚“为什么是 n 个根”",
         "“左图让 z 的辐角连续增加 2π 转 n 圈，右图让根的辐角只转 1 圈并依次扫过每个根。”",
         "AI 最初只让根点旋转，看不出与 z 的关系；按上述要求改成双面板对照后，动画才真正解释了多值性。"],
        ["5", "界面分区",
         "“窗口分四块：输入区、动画控制、绘图区、结果文本区；结果要同时给直角坐标和极坐标。”",
         "AI 生成的布局可用，我调整了控件顺序和默认尺寸，避免结果文本把画布挤小。"],
        ["6", "动图体积与画质",
         "“导出 GIF，尺寸 760x400 到 900x460，尽量控制在 3 MB 以内，文字不能糊。”",
         "AI 建议逐帧自适应调色板；我实测发现全局共享 200 色调色板体积更小且配色稳定，改用后者。"],
        ["7", "补充“两个复数相乘”的几何图像（评审意见）",
         "“增加两个复数相乘的几何演示：模长相乘、辐角相加，并说明它和 n 次方根的联系。”",
         "AI 提出用对数螺旋把乘法画成连续过程；我确认了 z1·z2^t 的模长与辐角公式后采纳，并追加“根自乘 n 次回到 z”的场景。"],
        ["8", "打包与文档",
         "“把这些文件放进一个 CFV 文件夹，写清运行方法和依赖。”",
         "AI 生成 README、run.bat、离线网页版；我核对了 run.bat 在缺少 Pillow 时的自动安装逻辑。"],
    ]

    debug_rows = [
        ["终端打印结果时报 UnicodeEncodeError",
         "Windows 控制台是 GBK 编码，负号“−”(U+2212) 不在 GBK 字符集里",
         "输出改用 ASCII 连字符“-”，并给标准输出加上 errors=\"replace\" 兜底",
         "在 PowerShell 与 cmd 下分别运行 --print，均不再报错"],
        ["输入 3+4i 直接报错",
         "Python 的 complex() 只认 3+4j 这种写法",
         "自写词法分析 + 安全 AST 求值，支持 i/j、π、省略乘号与 sqrt/exp/polar 等函数",
         "连续测试 15 种写法（含 sqrt(2)-sqrt(2)i、2(1+i)）全部解析正确"],
        ["z = -8 时辐角显示 -180°",
         "取负运算产生了 −0.0，atan2(−0, −8) 得到 −π",
         "在求辐角时先把 −0 归一化为 0",
         "z = -8 现在显示 180.00°，与数学习惯一致"],
        ["窗口切到后台再回来，动画会“跳”一大段",
         "两帧之间的时间差可能是几十秒",
         "把单帧时间差限制在 0.06 秒以内",
         "切后台 30 秒后回到窗口，画面仍连续"],
        ["GIF 体积接近 3 MB",
         "每帧使用各自的调色板，颜色表重复存储",
         "用第一帧量化出统一调色板，其余帧复用它（200 色）",
         "同一段动画由 2.65 MB 降到 2.07 MB，配色无跳变"],
        ["程序缩到很窄时标题文字出框",
         "标题与页脚是固定字号、固定文案",
         "按画布宽度分三档精简标题与页脚内容",
         "在 360×240 到 1400×700 之间共 126 帧渲染无出框"],
        ["在没装中文字体的机器上标签会变成方块",
         "PIL 默认字体不含中文字形",
         "启动时探测系统字体（msyh/simhei 等），找不到就自动切换英文标签",
         "把字体路径指向不存在的文件后，程序自动输出英文标签"],
        ["新增乘法场景时报 AttributeError: _draw_panel_back",
         "该方法原来只在 Scene 类里，新场景没有继承",
         "抽出 PanelScene 基类与模块级 draw_panel_back()，三种场景共用",
         "三个场景各渲染 4 种尺寸 × 多个参数，共 432 帧无异常"],
        ["GIF 导出 100 帧，读回来只有 49/81 帧",
         "PIL 会把连续相同的帧合并以减小体积（本程序末尾有停留帧）",
         "确认是预期行为：合并只影响帧数统计，画面与节奏不变",
         "逐帧比对第 0/30/60 帧内容不同，动画正常"],
        ["乘法场景里 t = 1 时乘积标记被移动点盖住",
         "绘制顺序是先画乘积、后画两条路径的动点",
         "调整为先画轨迹与动点，最后画乘积向量与标记",
         "像素校验：乘积点颜色为绿色 (126,231,135)"],
    ]

    review_rows = [
        ["未体现“以两个复数相乘图像为例子”",
         "新增“两个复数相乘”演示模式：z1·z2 的模长相乘、辐角相加，两条对数螺旋路径都终止于乘积；并增加“根自乘 n 次回到 z”的模式把乘法与开方连起来",
         "第 5、6 节，附录 A 的运行示例"],
        ["材料过于单一，缺少过程记录",
         "本报告补充了 AI 提示词迭代记录（8 轮）与调试记录（10 条），并列出每个问题的现象、原因、处理与验证方式",
         "第 7、8 节"],
        ["缺少对几何意义的说明与个人反思",
         "第 2 节给出完整推导与几何解释，第 9 节记录我个人理解的转变、验证方法上的收获和不足",
         "第 2、9 节"],
        ["工程规范上的小瑕疵：颜色常量定义在类之后",
         "把 RGB 与 #rrggbb 两套颜色常量集中到文件开头的配色区，并抽出 PanelScene 基类消除重复的绘制代码",
         "complex_roots.py 顶部配色区"],
        ["缺少独立的说明文档",
         "CFV 目录内含 README.md（用法、参数表、复数写法、常见问题）与本实验报告",
         "CFV/README.md、本报告"],
    ]

    file_rows = [
        ["complex_roots.py", "主程序：三种演示模式 + 交互界面 + 动图/图片导出"],
        ["run.bat", "双击启动，自动选择带 tkinter 的 Python 并补齐 Pillow"],
        ["complex_roots_web.html", "离线网页版，浏览器打开即可输入复数查看结果"],
        ["web_fragment.html / build_web.py", "网页版源片段与重新生成脚本"],
        ["roots_demo.gif / mul_demo.gif / power_demo.gif", "三段示例动图（根分布、两数相乘、根自乘）"],
        ["report_figures/", "本报告使用的配图与公式图片"],
        ["build_report.py", "本报告的生成脚本"],
        ["README.md / requirements.txt", "说明文档与依赖清单"],
    ]

    input_samples = ["3+4i", "-8", "2i", "sqrt(2)-sqrt(2)i",
                     "2*exp(i*pi/3)", "polar(2, pi/6)", "2(1+i)"]
    input_notes = {
        "3+4i": "直角坐标，虚数单位可写 i 或 j",
        "-8": "负实数：辐角为 180°",
        "2i": "纯虚数：省略实部",
        "sqrt(2)-sqrt(2)i": "带根号与省略乘号",
        "2*exp(i*pi/3)": "指数形式，π 与 i 都能识别",
        "polar(2, pi/6)": "极坐标形式：模长 2、辐角 π/6",
        "2(1+i)": "省略乘号的两个括号",
    }
    input_rows = [[s, cr.fmt_complex(cr.parse_complex(s), 6), input_notes[s]]
                  for s in input_samples]

    # ---- 开始写文档 ----
    doc = Document()
    setup_styles(doc)
    add_page_number_footer(doc)

    doc.add_paragraph("复数 n 次方根可视化工具 实验报告", style="Title")
    para(doc, "Python 实现，含复数乘法的几何演示与交互式小程序", size=11, cn="微软雅黑",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, color=GRAY)
    add_table(doc, ["项目", "内容"], [
        ["课程名称", "　"],
        ["姓名 / 学号", "　"],
        ["提交日期", "2026 年 9 月 18 日"],
        ["作品文件", "complex_roots.py（Python 3.8+，依赖 Pillow）"],
    ], widths=[3.2, 12.8], align_center_cols={0, 1})

    para(doc, "摘要", size=11, bold=True, cn="微软雅黑", space_after=4)
    para(doc, "本报告说明我用 Python 完成的复数 n 次方根可视化工具。程序接受任意复数 z 与方根次数 n，"
              "按德莫弗定理给出全部 n 个根（直角坐标与极坐标），并用动图说明这些根的来源；"
              "同时提供“两个复数相乘”的几何演示，说明模长相乘、辐角相加的规律，以及它与 n 次方根之间的关系。"
              "程序是单文件实现，只依赖标准库和 Pillow，另外提供一个可离线打开的网页版本。"
              "报告记录了数学推导、程序结构与关键代码、AI 提示词的迭代过程、调试中遇到的十个问题及其处理办法，"
              "最后是我对这次实践的反思与后续改进方向。")

    # ================= 一 =================
    h1(doc, "一、任务理解与完成情况")
    para(doc, "任务要求生成复数 n 次方根的动图，并进一步做成可以随时输入复数和方根次数的交互式小程序；"
              "任务中还特别提出“以两个复数相乘图像为例子”。我把这三点理解为一条完整的线索："
              "复数乘法的几何图像是“伸缩 + 旋转”，而求 n 次方根正是反过来问“哪个复数自乘 n 次等于 z”，"
              "所以程序既要能算、能画，也要把乘法与开方的关系演示出来。")
    para(doc, "为此，程序提供三种演示模式：n 次方根的分布与辐角扫过、两个复数相乘、某个根自乘 n 次回到 z；"
              "三种模式共用同一套画法与交互界面，都能导出 GIF 动图。表 1 对照了任务要求与实现情况。")
    table_caption(doc, "表 1　任务要求与实现对照")
    add_table(doc, ["任务要求", "本作品的实现", "位置"], [
        ["生成复数 n 次方根的动图", "三种模式均可导出 GIF：根分布、两数相乘、根自乘", "第 5、6 节"],
        ["以两个复数相乘图像为例子", "专门的乘法场景：模长相乘、辐角相加，两条对数螺旋路径", "第 5 节"],
        ["任意复数与方根次数的交互式小程序", "tkinter 三模式界面，支持 3+4i、sqrt(2)-sqrt(2)i、2*exp(i*pi/3) 等写法", "第 3、4 节"],
        ["结果与图形同时输出", "结果面板给出直角坐标、极坐标与 |w_k^n − z| 自校验值", "第 6 节"],
    ], widths=[4.6, 8.4, 3.0], align_center_cols={2})

    # ================= 二 =================
    h1(doc, "二、数学原理")
    h2(doc, "2.1 复数乘法：模长相乘、辐角相加")
    para(doc, "把两个复数写成三角形式 z1 = r1(cos θ1 + i sin θ1)，z2 = r2(cos θ2 + i sin θ2)，"
              "相乘后展开并合并同类项，利用和角公式即可得到乘法的三角形式，再由欧拉公式写成指数形式：")
    add_equation(doc, f_mul1)
    add_equation(doc, f_mul2)
    para(doc, "两式合起来给出乘法最常用的两条结论：积的模长等于两个模长之积，积的辐角等于两个辐角之和。"
              "几何上，把 z1 的模长拉伸（或压缩）|z2| 倍，再逆时针旋转 θ2，终点就是 z1·z2；"
              "反过来，把 z2 拉伸 |z1| 倍、旋转 θ1，得到的仍是同一点。")

    h2(doc, "2.2 对数螺旋：把乘法画成连续过程")
    para(doc, "程序里没有把乘法画成“一步跳过去”，而是引入连续参数 t ∈ [0, 1]，让中间量沿 z1·z2^t 移动。"
              "它的模长是 r1r2^t，辐角是 θ1 + tθ2，两者同时变化，因此轨迹是一条对数螺旋：")
    add_equation(doc, f_spiral)
    para(doc, "t = 0 时在 z1，t = 1 时在 z1·z2。旋转与伸缩被拆成可以看得见的两个动作，"
              "这也让“辐角相加”不再是一句口诀，而是画面上指针转过的角度。")

    h2(doc, "2.3 德莫弗定理与 n 次方根")
    para(doc, "在乘法结论中取 n 个相同因子，就得到德莫弗定理：")
    add_equation(doc, f_demoivre)
    para(doc, "设 z = r·e^(iθ)，若 w^n = z，把 w 写成 ρ·e^(iφ)，则 ρ^n = r 且 nφ = θ + 2πk（k 为整数）。"
              "ρ 只能取非负的 r^(1/n)，而 φ 在模 2π 意义下都对应同一个 w，于是得到 n 个不同的根：")
    add_equation(doc, f_root)
    para(doc, "这 n 个根均匀分布在以原点为圆心、半径 r^(1/n) 的圆上，相邻夹角为 2π/n。"
              "程序中的动画正是围绕这一点设计的：z 的辐角每加 2π 仍然是同一个 z，"
              "而把 θ + 2πk 除以 n 之后，落在了 n 个互不相同的方向上。")

    h2(doc, "2.4 乘法与开方的联系")
    para(doc, "把“求 n 次方根”与“两个复数相乘”放在一起看，任务里提到的乘法图像就自然接上了："
              "根就是那个自乘 n 次等于 z 的数。每乘一次 w，模长乘以 |w|、辐角加上 arg w，n 次之后必须回到 z：")
    add_equation(doc, f_check)
    para(doc, "程序中的“根自乘 n 次”模式把这条链画成螺旋：第 j 步落在 w^j 上，第 n 步正好落在 z 上。"
              "这样一来，乘法不再只是公式里的展开，而是开方运算的逆过程。")

    # ================= 三 =================
    h1(doc, "三、程序总体设计")
    para(doc, "程序是一个单文件 Python 程序，按“输入解析 → 数学计算 → 场景渲染 → 交互与导出”四层组织。"
              "三种演示模式各自是一个场景类，但它们共用同一个面板基类，改一处网格或坐标轴的画法，三个模式同时生效。")
    table_caption(doc, "表 2　程序的主要模块")
    add_table(doc, ["模块 / 类", "职责"], [
        ["parse_complex()", "把 3+4i、sqrt(2)-sqrt(2)i、2*exp(i*pi/3) 等文本解析为复数（词法分析 + 安全 AST 求值）"],
        ["nth_roots_of()", "按德莫弗定式计算 n 个根，返回根、|z|、arg z 与 r^(1/n)"],
        ["PanelScene / draw_panel_back()", "三种场景共用的面板底图：网格、刻度、坐标轴、单位圆"],
        ["Scene", "n 次方根场景：左图看 z 的辐角，右图看 n 个根与等角分布"],
        ["MultiplyScene", "两个复数相乘：两条对数螺旋路径、角度相加的扇形"],
        ["PowerScene", "根自乘 n 次回到 z：螺旋、整数次幂标记与“连乘账本”"],
        ["export_gif() / export_scene_gif() 等", "把任意场景按统一进度渲染成帧并导出 GIF 或 PNG"],
        ["RootsApp", "tkinter 交互界面：模式切换、参数输入、动画控制、结果面板与导出"],
        ["complex_roots_web.html", "离线网页版，用同一套数学与绘制逻辑在浏览器中运行"],
    ], widths=[4.6, 11.4])
    para(doc, "界面按用途分成四块：顶部是模式与参数输入（复数 z、方根次数 n，乘法模式下换成第二个复数 z2，"
              "自乘模式下选择第几个根 k）；第二行是动画控制与显示开关；左侧是绘图区，右侧是结果面板；"
              "底部是导出与复制按钮。绘图区会随窗口缩放自动重排，窄窗口下自动精简标题文字。")

    # ================= 四 =================
    h1(doc, "四、关键实现")
    h2(doc, "4.1 n 次方根的计算与自校验")
    para(doc, "核心计算只有几行，但把“多值性”写进了下标：每个 k 对应一个 2πk 的偏移量。")
    code_block(doc, snippet("def nth_roots_of", 11))
    para(doc, "结果面板最后一行给出 max |w_k^n − z|，用来说明这 n 个数确实都是根。"
              "以 z = 3+4i、n = 5 为例，这个误差在 10⁻¹⁵ 量级，属于浮点舍入的正常范围。")

    h2(doc, "4.2 复数输入的解析")
    para(doc, "直接调用 complex() 只能识别 3+4j 这种写法，无法满足“随时输入一个复数”的要求。"
              "程序先用正则把文本切成记号，再在两类记号之间补上省略的乘号，最后交给受限的 AST 求值器计算。")
    code_block(doc, snippet("# 补上省略的乘号", 16, back=2))

    h2(doc, "4.3 三种场景共用一套画法")
    para(doc, "新增乘法与自乘场景时，如果各自复制一份网格与坐标轴的绘制代码，后期改动很容易漏改。"
              "我把面板底图抽成模块级函数，并让三个场景继承同一个基类：")
    code_block(doc, snippet("def layout_panels", 12))
    code_block(doc, snippet("class PanelScene", 8))

    h2(doc, "4.4 动图导出")
    para(doc, "三种场景都实现 render_progress(u)，把 0 到 1 的进度映射到自己的动画参数，"
              "因此同一套导出流程可以处理全部模式。导出时用首帧生成统一调色板，其余帧复用它，"
              "既压缩了体积，也避免颜色在帧之间跳动。")
    code_block(doc, snippet("def save_gif", 8))

    # ================= 五 =================
    h1(doc, "五、两个复数相乘的几何演示")
    para(doc, "这是本次修订补充的重点，对应任务中“以两个复数相乘图像为例子”的要求。"
              "以 z1 = 3 + i、z2 = 2 + i 为例：|z1| = 3.1623，θ1 = 18.43°，|z2| = 2.2361，θ2 = 26.57°。"
              "程序同时画出两条路径：把 z1 伸缩 |z2| 倍并旋转 θ2，或者把 z2 伸缩 |z1| 倍并旋转 θ1，"
              "两条螺旋最终落在同一点。")
    add_figure(doc, os.path.join(FIG_DIR, "fig2_multiply.png"),
               "图 1　两个复数相乘（z1 = 3 + i，z2 = 2 + i）：左图为两个因数与它们各自的辐角，"
               "右图的两条螺旋分别从 z1、z2 出发，终止于同一点 z1·z2 = 5 + 5i", 15.0)
    para(doc, "图 1 中右侧的粉色扇形是角度相加的过程：先转过 θ1，再接着转过 θ2，"
              "两次旋转拼起来正好是积的辐角。图中的绿色虚线圆半径就是积的模长，数值上与 |z1|·|z2| 一致。")
    add_figure(doc, os.path.join(FIG_DIR, "fig3_multiply_steps.png"),
               "图 2　乘法动画的两个阶段：t = 0.25 时模长与辐角同时变化，t = 1.00 时到达 z1·z2", 8.4)
    table_caption(doc, "表 3　两条乘法路径（z1 = 3 + i，z2 = 2 + i）")
    add_table(doc, ["路径", "操作", "终点"], mul_rows,
              widths=[2.8, 9.4, 3.8], align_center_cols={0, 2})
    para(doc, "这段演示解释了乘法，也解释了为什么“开方”必须回到乘法："
              "把 w 连续乘 n 次，就是让模长连续乘以 |w|、辐角连续加上 arg w，"
              "落到 z 上的那个 w 就是 z 的 n 次方根。下一节的“根自乘 n 次”模式把这条链条完整画了出来。")

    # ================= 六 =================
    h1(doc, "六、n 次方根的动图与自乘链条")
    para(doc, "n 次方根模式下，左侧面板显示原复数 z 与它的辐角，右侧面板显示 n 个根、"
              "外接圆与正 n 边形。动画让 z 的辐角连续增加，转满 n 圈；同一时间根的辐角只转一圈，"
              "依次扫过 n 个根，屏幕下方的数值同步显示当前扫到哪一个根。")
    add_figure(doc, os.path.join(FIG_DIR, "fig1_roots.png"),
               "图 3　n 次方根场景（z = 3 + 4i，n = 5）：左图为原复数与辐角，"
               "右图 5 个根等角分布在半径 |z|^(1/5) 的圆上，相邻夹角 72°", 15.0)
    para(doc, "“根自乘 n 次”模式把 w = w0 的 1 到 n 次幂依次画在螺旋上，"
              "每一步的角度扇形都是同一次乘法；右侧的“连乘账本”列出每一步的数值，最后一行回到 z。")
    add_figure(doc, os.path.join(FIG_DIR, "fig4_power.png"),
               "图 4　根自乘 n 次回到 z（z = 3 + 4i，n = 5，取 k = 0）："
               "螺旋上标出 w^1 到 w^5，第 5 步落在 z 上", 15.0)
    table_caption(doc, "表 4　w 的连乘账本（z = 3 + 4i，n = 5，取 k = 0）")
    add_table(doc, ["步数 j", "w^j", "模长 |w|^j", "辐角", "说明"], power_rows,
              widths=[2.0, 5.4, 3.2, 3.0, 2.4], align_center_cols={0, 2, 3, 4})
    para(doc, "表中模长按 |w|^j 递增，辐角每次增加同样的 10.63°，第五步模长回到 |z| = 5、"
              "辐角回到 θ + 2πk，所以落点正是 z。这张表是把第 2.4 节的公式代入具体数字后的结果。")

    # ================= 七 =================
    h1(doc, "七、AI 提示词迭代记录")
    para(doc, "开发过程中我主要把 AI 当作实现助手：数学结论由我给出，代码结构由我约束，"
              "AI 负责把要求写成可运行的代码，我再逐条验证。表 5 记录了关键轮次的提示词要点与我的处理。")
    table_caption(doc, "表 5　AI 提示词迭代记录")
    add_table(doc, ["轮次", "我的目标", "关键提示词（节选）", "AI 输出与我的处理"], prompt_rows,
              widths=[1.3, 3.6, 5.6, 5.5], align_center_cols={0})
    para(doc, "几轮下来最直接的体会是：提示词里必须写清约束条件。早期只说“画个动图”，"
              "AI 会用最常见的 matplotlib + numpy；把“环境里没有这两个库、只能用标准库和 Pillow”写进去之后，"
              "方案立刻落到可运行的范围内。同样，把“非法输入不能崩”“体积控制在 3 MB 以内”这类可检验的要求写进提示词，"
              "得到的结果才容易验收。")

    # ================= 八 =================
    h1(doc, "八、调试过程与问题记录")
    para(doc, "表 6 是我在开发和修订过程中实际遇到的问题。每条都记录了现象、原因、处理方式和验证方法，"
              "其中绝大多数问题不是“写不出来”，而是“写出来但没考虑边界”。")
    table_caption(doc, "表 6　调试过程与问题记录")
    add_table(doc, ["现象", "原因", "处理", "验证方式"], debug_rows,
              widths=[3.5, 4.2, 4.6, 3.7], size=9)

    # ================= 九 =================
    h1(doc, "九、个人反思")
    para(doc, "做完这次实践，我对“复数乘法的几何意义”的理解比做题时清楚得多。"
              "以前提到德莫弗定理，我记住的是公式 (cos θ + i sin θ)^n = cos nθ + i sin nθ，"
              "把它当成三角恒等式来背。为了画乘法动画，我必须先想清楚 z1·z2^t 的模长和辐角随 t 怎么变，"
              "才写出让点沿螺旋前进的代码，这一步逼着我承认：乘法就是缩放与旋转的合成，"
              "而开方就是把这个过程倒过来问“谁自乘 n 次能得到它”。把两种演示放在同一个界面里对照，"
              "我才真正接受了“n 次方根有 n 个”不是因为公式里多了个 k，而是因为辐角可以在模 2π 意义上绕 n 次。")
    para(doc, "第二个体会是验证方法比“看着像对”更重要。程序画出来的图很直观，但直观不等于正确。"
              "我给每个根加了 |w_k^n − z| 的误差输出，又用像素位置反查：把公式算出的坐标换算成像素，"
              "检查那个位置的颜色是不是根的颜色、产品点是不是落在螺旋的终点上。"
              "几次“看起来没问题”的地方正是这样被查出来的，比如 -8 的辐角显示成 −180°、"
              "乘法场景里乘积标记被移动点盖住。")
    para(doc, "第三个体会关于 AI 的使用方式。AI 能很快给出结构完整的代码，但它默认的技术选型、"
              "默认的容错水平、默认的界面布局都未必符合我的环境与要求；能省下的是打字的时间，"
              "省不下的是把问题说清楚和把结果检查一遍的时间。我在提示词里逐步补上“依赖限制、"
              "输入格式、数值验证、体积上限”这些约束之后，返工的次数明显减少。")
    para(doc, "不足也很明显。解析器目前只支持字面量与白名单函数，不能定义变量或保存历史输入；"
              "动图仍是 2D 的，没有办法显示模长与辐角分开看的三维效果；"
              "工程上还缺少自动化的单元测试，验证主要靠我手动运行与像素比对，"
              "颜色常量放在类之后这类不规范的地方也是在评审指出后才集中整理。"
              "下一步我打算给核心函数补上 pytest 测试，并把解析器的错误提示做得更具体。")

    # ================= 十 =================
    h1(doc, "十、评审意见对照")
    table_caption(doc, "表 7　评审意见与本版改进")
    add_table(doc, ["评审意见", "本版改进", "位置"], review_rows, widths=[4.2, 9.3, 2.5])

    # ================= 附录 =================
    h1(doc, "附录 A　运行方式与文件清单")
    para(doc, "程序只需要 Python 3.8 以上和 Pillow（图形界面额外需要 tkinter，Windows 官方安装包自带）。"
              "双击 run.bat 会打开交互窗口；命令行方式支持三种演示模式：")
    code_block(doc, "\n".join([
        "python complex_roots.py                                  # 打开交互窗口（默认 n 次方根模式）",
        "python complex_roots.py -z \"3+4i\" -n 5 --print            # 打印 5 个根",
        "python complex_roots.py --mode multiply -z \"3+1i\" --z2 \"2+1i\" --save mul.gif",
        "python complex_roots.py --mode power -z \"3+4i\" -n 5 -k 0 --save power.gif",
    ]), size=9)
    table_caption(doc, "表 8　文件清单")
    add_table(doc, ["文件", "说明"], file_rows, widths=[6.0, 10.0])

    h1(doc, "附录 B　演示模式与输入写法速查")
    table_caption(doc, "表 9　三种演示模式")
    add_table(doc, ["模式", "需要的输入", "画面内容", "对应命令"], [
        ["n 次方根（默认）", "复数 z、方根次数 n",
         "左图 z 与辐角，右图 n 个根与正 n 边形，指针依次扫过每个根",
         "python complex_roots.py -z \"3+4i\" -n 5"],
        ["两个复数相乘", "复数 z1、z2",
         "左图两个因数与辐角，右图两条对数螺旋与角度相加的扇形",
         "python complex_roots.py --mode multiply -z \"3+1i\" --z2 \"2+1i\""],
        ["根自乘 n 次", "复数 z、次数 n、根的序号 k",
         "螺旋上标出 w^1 到 w^n，右侧列出连乘账本",
         "python complex_roots.py --mode power -z \"3+4i\" -n 5 -k 0"],
    ], widths=[2.6, 3.0, 6.3, 3.6], size=9)

    table_caption(doc, "表 10　复数输入的写法与解析结果")
    add_table(doc, ["输入写法", "解析出的复数", "说明"], input_rows,
              widths=[4.2, 5.0, 6.3], size=9)
    para(doc, "表中结果由程序现算，与 parse_complex() 一致；输入非法表达式时会在状态栏提示错误，不会崩溃。")

    doc.save(OUT_PATH)
    print("已生成:", OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    sys.exit(0 if build() else 1)
