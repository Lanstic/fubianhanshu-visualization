#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 complex_roots.py 输出成一份便于提交的源码清单 Word 文档。

用法：python build_code_doc.py
输出：源码code（更新版）.docx
"""

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "complex_roots.py")
OUT = os.path.join(HERE, "源码code（更新版）.docx")

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x59, 0x59, 0x59)


def set_font(run, cn, latin, size, bold=False, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def main():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = doc.styles["Title"]
    title.font.name = "微软雅黑"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.paragraph_format.space_after = Pt(2)
    ppr = title.element.get_or_add_pPr()
    for bdr in ppr.findall(qn("w:pBdr")):          # 去掉标题自带的下框线
        ppr.remove(bdr)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"complex_roots.py　共 {len(lines)} 行　Python 3.8+　依赖：Pillow")
    set_font(run, "微软雅黑", "Times New Roman", 9.5, color=GRAY)

    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.03
        run = para.add_run(line if line.strip() else " ")
        set_font(run, "Consolas", "Consolas", 7.5, color=RGBColor(0x1F, 0x2A, 0x37))

    doc.save(OUT)
    print("已生成:", OUT, "（", len(lines), "行 ）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
